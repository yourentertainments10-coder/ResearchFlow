"""Track A: loading CSV, XLSX and JSON, with the limits and hostile inputs of design section 8."""

from __future__ import annotations

import datetime as dt
import io
import re
import time
import zipfile

import pandas as pd
import pytest

from sie.datasets import from_frame, open_dataset
from sie.datasets.models import (
    DatasetError,
    LimitExceeded,
    LoadOptions,
    ParseError,
)
from sie.datasets.models import Limits as L
from sie.datasets.provenance import sha256_bytes


def csv_ds(text: str | bytes, **opts):
    data = text.encode() if isinstance(text, str) else text
    return open_dataset(data, "t.csv", LoadOptions(**opts) if opts else None)


def note(ds, code):
    found = [n for n in ds.notes if n["code"] == code]
    return found[0] if found else None


# --- CSV -----------------------------------------------------------------------------------------


def test_csv_basic_row_ids_empty_cells_and_source_reference():
    data = b"name,score\nAnn,10\nBob,\n"
    ds = open_dataset(data, "scores.csv", kind="upload", retrieved_at="2026-10-13T00:00:00Z")
    assert ds.columns == ["name", "score"]
    assert ds.frame.index.tolist() == [1, 2] and ds.frame.index.name == "row_id"
    assert ds.frame.loc[1, "score"] == "10" and ds.frame.loc[2, "score"] is None
    assert str(ds.frame["name"].dtype) == "object"
    assert ds.source.sha256 == sha256_bytes(data) and ds.source.size_bytes == len(data)
    assert ds.source.name == "scores.csv" and ds.source.retrieved_at == "2026-10-13T00:00:00Z"
    assert ds.rejects is None


def test_csv_utf8_bom_is_removed_and_other_unicode_survives():
    ds = csv_ds("﻿city,x\nZürich,1\n")
    assert ds.columns == ["city", "x"] and ds.frame.loc[1, "city"] == "Zürich"


def test_csv_quoted_fields_with_delimiters_and_newlines():
    ds = csv_ds('a,b\n"x,1","line1\nline2"\n"he said ""hi""",2\n')
    assert ds.frame.loc[1, "a"] == "x,1" and ds.frame.loc[1, "b"] == "line1\nline2"
    assert ds.frame.loc[2, "a"] == 'he said "hi"'


def test_csv_cells_starting_with_equals_are_plain_text():
    ds = csv_ds('a\n=1+1\n=HYPERLINK("http://x")\n')
    assert ds.frame["a"].tolist() == ["=1+1", '=HYPERLINK("http://x")']


def test_csv_header_only_has_zero_rows_and_columns():
    ds = csv_ds("a,b\n")
    assert ds.rows == 0 and ds.columns == ["a", "b"]


def test_csv_blank_lines_are_skipped_but_row_ids_still_count_them():
    ds = csv_ds("a\n1\n\n3\n")
    assert ds.frame.index.tolist() == [1, 3]
    assert note(ds, "blank_lines_skipped")["count"] == 1


def test_csv_empty_file_and_no_header_are_errors():
    with pytest.raises(ParseError, match="empty"):
        csv_ds("")
    with pytest.raises(ParseError, match="empty"):
        csv_ds("\n\n")


def test_csv_ragged_rows_fail_with_row_and_counts_by_default():
    with pytest.raises(ParseError) as info:
        csv_ds("a,b\n1,2\n3\n")
    assert "row 2" in str(info.value) and "1 fields" in str(info.value) and "2" in str(info.value)


def test_csv_ragged_reject_sets_rows_aside_with_reason():
    ds = csv_ds("a,b\n1,2\n3\n4,5,6\n7,8\n", ragged="reject")
    assert ds.frame.index.tolist() == [1, 4]
    assert ds.rejects is not None
    assert ds.rejects["row_id"].tolist() == [2, 3]
    assert set(ds.rejects["reason"]) == {"RAGGED_ROW"}
    assert list(ds.rejects.columns) == ["row_id", "operation", "column", "original_value", "reason"]
    assert note(ds, "ragged_rows_rejected")["count"] == 2


def test_csv_duplicate_and_empty_headers_are_errors_naming_positions():
    with pytest.raises(ParseError) as info:
        csv_ds("a,b,a,\n1,2,3,4\n")
    message = str(info.value)
    assert "'a' at positions [1, 3]" in message and "positions [4]" in message


def test_csv_headers_compare_after_trimming():
    with pytest.raises(ParseError, match="duplicate"):
        csv_ds("a, a \n1,2\n")


def test_csv_rename_duplicates_is_deterministic_and_logged():
    ds = csv_ds("a,a,a__2,,a\n1,2,3,4,5\n", rename_duplicates=True)
    assert ds.columns == ["a", "a__3", "a__2", "column_4", "a__4"]
    renames = note(ds, "renamed_header")
    assert renames["count"] == 3
    assert {r["to"] for r in renames["detail"]} >= {"a__3", "column_4", "a__4"}
    again = csv_ds("a,a,a__2,,a\n1,2,3,4,5\n", rename_duplicates=True)
    assert again.columns == ds.columns


def test_csv_undecodable_bytes_report_the_offset():
    with pytest.raises(ParseError) as info:
        csv_ds(b"a\n\xe9\n")
    assert "byte offset 2" in str(info.value)


def test_csv_other_encoding_must_be_stated():
    data = "a\ncafé\n".encode("latin-1")
    assert csv_ds(data, encoding="latin-1").frame.loc[1, "a"] == "café"
    with pytest.raises(DatasetError, match="unknown encoding"):
        csv_ds(b"a\n1\n", encoding="no-such-codec")


def test_csv_nul_bytes_are_an_error():
    with pytest.raises(ParseError, match="NUL"):
        csv_ds(b"a,b\n1,\x002\n")


def test_csv_delimiter_is_never_guessed():
    with pytest.raises(ParseError, match="semicolon"):
        csv_ds("a;b\n1;2\n")
    with pytest.raises(ParseError, match="tab"):
        csv_ds("a\tb\n1\t2\n")
    assert csv_ds("a;b\n1;2\n", delimiter=";").columns == ["a", "b"]
    assert csv_ds("a\tb\n1\t2\n", delimiter="\t").frame.loc[1, "b"] == "2"


def test_csv_malformed_quote_is_an_error_with_the_line():
    with pytest.raises(ParseError, match="line 2"):
        csv_ds('a,b\n"x"y,2\n')


def test_csv_leading_zeros_and_number_like_text_are_not_changed_at_load():
    ds = csv_ds("id,v\n007,1,5\n".replace("1,5", '"1,5"'))
    assert ds.frame.loc[1, "id"] == "007" and ds.frame.loc[1, "v"] == "1,5"


# --- limits --------------------------------------------------------------------------------------


def test_limit_file_size_is_checked_before_parsing():
    with pytest.raises(LimitExceeded) as info:
        open_dataset(b"a\n" + b"1\n" * 20, "t.csv", limits=L(max_bytes=10))
    assert info.value.limit == "max_bytes"


@pytest.mark.parametrize(
    ("text", "limits", "name"),
    [
        ("a\n1\n2\n3\n", L(max_rows=2), "max_rows"),
        ("a,b,c\n1,2,3\n", L(max_columns=2), "max_columns"),
        ("a,b\n1,2\n3,4\n", L(max_cells=3), "max_cells"),
        ("a\n" + "x" * 11 + "\n", L(max_cell_chars=10), "max_cell_chars"),
    ],
)
def test_csv_limits_name_the_limit_that_triggered(text, limits, name):
    with pytest.raises(LimitExceeded) as info:
        open_dataset(text.encode(), "t.csv", limits=limits)
    assert info.value.limit == name


def test_csv_rows_exactly_at_the_limit_are_accepted():
    assert open_dataset(b"a\n1\n2\n", "t.csv", limits=L(max_rows=2)).rows == 2


def test_csv_a_field_larger_than_the_parser_allows_is_a_limit_not_a_crash():
    big = b"a\n" + b"x" * 200_000 + b"\n"
    with pytest.raises(LimitExceeded) as info:
        open_dataset(big, "t.csv", limits=L(max_cell_chars=1_000_000))
    assert info.value.limit == "max_cell_chars"


def test_parse_time_budget_is_enforced():
    data = b"a\n" + b"1\n" * 10_050
    started = time.monotonic()
    with pytest.raises(LimitExceeded) as info:
        open_dataset(data, "t.csv", limits=L(parse_seconds=-1))
    assert info.value.limit == "parse_seconds" and time.monotonic() - started < 5


def test_unsupported_and_legacy_extensions_are_rejected():
    for name in ("a.xls", "a.xlsb", "a.xlsm"):
        with pytest.raises(DatasetError, match="not supported"):
            open_dataset(b"x", name)
    for name in ("a.txt", "noext", "a.csv.exe"):
        with pytest.raises(DatasetError, match="unsupported file type"):
            open_dataset(b"x", name)
    assert open_dataset(b"a\n1\n", "UPPER.CSV").rows == 1


# --- JSON ----------------------------------------------------------------------------------------


def js(text: str | bytes, **opts):
    data = text.encode() if isinstance(text, str) else text
    return open_dataset(data, "t.json", LoadOptions(**opts) if opts else None)


def test_json_basic_types_are_kept_as_the_source_typed_them():
    ds = js('[{"a": 1, "b": 1.0, "c": "x", "d": true, "e": null}]')
    row = ds.frame.loc[1]
    assert row["a"] == 1 and isinstance(row["a"], int) and not isinstance(row["a"], bool)
    assert isinstance(row["b"], float) and row["c"] == "x" and row["d"] is True and row["e"] is None
    assert ds.frame.index.tolist() == [1]


def test_json_columns_are_the_union_in_order_and_missing_keys_are_counted():
    ds = js('[{"a": 1}, {"b": 2, "a": 3}, {"c": 4}]')
    assert ds.columns == ["a", "b", "c"]
    assert ds.frame["b"].tolist() == [None, 2, None]
    assert note(ds, "missing_keys")["count"] == 5


def test_json_duplicate_keys_are_an_error():
    with pytest.raises(ParseError, match="duplicate key 'a'"):
        js('[{"a": 1, "a": 2}]')


@pytest.mark.parametrize("token", ["NaN", "Infinity", "-Infinity", "1e999"])
def test_json_non_finite_numbers_are_rejected(token):
    with pytest.raises(ParseError):
        js(f'[{{"a": {token}}}]')


def test_json_integers_beyond_int64_are_an_error_not_a_float():
    with pytest.raises(ParseError, match="64-bit"):
        js('[{"a": 9223372036854775808}]')
    assert js('[{"a": 9223372036854775807}]').frame.loc[1, "a"] == 9223372036854775807


def test_json_malformed_reports_line_and_column_and_trailing_text_is_an_error():
    with pytest.raises(ParseError, match=r"line 2, column"):
        js('[\n{"a": }]')
    with pytest.raises(ParseError, match="Extra data"):
        js('[{"a": 1}] garbage')


def test_json_must_be_an_array_of_objects():
    with pytest.raises(ParseError, match="array of objects"):
        js('{"a": 1}')
    with pytest.raises(ParseError, match="row 2"):
        js('[{"a": 1}, 5]')
    with pytest.raises(ParseError, match="empty"):
        js("[]")


def test_json_nested_values_are_rejected_with_their_path():
    with pytest.raises(ParseError) as info:
        js('[{"a": 1}, {"a": {"b": 2}}]')
    assert "row 2" in str(info.value) and "$[1].a" in str(info.value)
    with pytest.raises(ParseError, match="array"):
        js('[{"a": [1]}]')


def test_json_depth_is_limited_and_deep_documents_do_not_crash():
    deep_value = "[" * 70 + "]" * 70
    with pytest.raises(LimitExceeded) as info:
        js(f'[{{"a": {deep_value}}}]')
    assert info.value.limit == "max_json_depth"
    hostile = "[" * 200_000
    with pytest.raises((LimitExceeded, ParseError)):
        js(hostile)


def test_json_bom_and_unicode_and_row_limit():
    assert js("﻿" + '[{"n": "é"}]').frame.loc[1, "n"] == "é"
    with pytest.raises(LimitExceeded) as info:
        open_dataset(b'[{"a":1},{"a":2},{"a":3}]', "t.json", limits=L(max_rows=2))
    assert info.value.limit == "max_rows"


def test_json_duplicate_names_after_trimming_follow_the_header_policy():
    with pytest.raises(ParseError, match="duplicate"):
        js('[{"a": 1, " a": 2}]')
    assert js('[{"a": 1, " a": 2}]', rename_duplicates=True).columns == ["a", "a__2"]


# --- XLSX ----------------------------------------------------------------------------------------


def make_xlsx(fill, *, hidden_sheet: bool = False) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Data"
    fill(sheet, book)
    if hidden_sheet:
        extra = book.create_sheet("Secret")
        extra["A1"] = "s"
        extra.sheet_state = "hidden"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def rewrite(data: bytes, edit) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            payload = edit(info.filename, payload)
            target.writestr(info.filename, payload)
    return out.getvalue()


def xl(data: bytes, **opts):
    return open_dataset(data, "t.xlsx", LoadOptions(**opts) if opts else None)


def simple(sheet, _book):
    sheet.append(["name", "n", "when", "ok"])
    sheet.append(["Ann", 3, dt.datetime(2026, 1, 2), True])
    sheet.append(["Bob", 1.5, None, False])


def test_xlsx_values_keep_their_native_types_and_row_ids():
    ds = xl(make_xlsx(simple))
    assert ds.columns == ["name", "n", "when", "ok"]
    assert ds.frame.index.tolist() == [1, 2]
    assert ds.frame.loc[1, "n"] == 3 and ds.frame.loc[2, "n"] == 1.5
    assert ds.frame.loc[1, "when"] == dt.datetime(2026, 1, 2) and ds.frame.loc[2, "when"] is None
    assert ds.frame.loc[1, "ok"] is True
    assert note(ds, "excel_epoch")["count"] == 1900


def test_xlsx_formulas_are_never_evaluated_cached_value_is_read():
    def fill(sheet, _book):
        sheet.append(["x", "y"])
        sheet.append([1, "=A2*2"])
        sheet.append([2, "=A3*2"])

    data = make_xlsx(fill)

    def add_cache(name, payload):
        if name == "xl/worksheets/sheet1.xml":
            text = payload.decode()
            text = re.sub(
                r"(<f>A2\*2</f>)<v\s*/>|(<f>A2\*2</f>)<v></v>", r"\1\2<v>2</v>", text, count=1
            )
            return text.encode()
        return payload

    ds = xl(rewrite(data, add_cache))
    assert ds.frame["y"].tolist()[0] == 2  # the stored value, not computed
    assert ds.frame["y"].tolist()[1] is None  # no cached value: null, never computed
    assert note(ds, "formula_cells")["count"] == 2
    assert note(ds, "formula_without_cached_value")["count"] == 1


def test_xlsx_formula_text_never_appears_as_a_value():
    def fill(sheet, _book):
        sheet.append(["x"])
        sheet.append(['=HYPERLINK("http://evil","x")'])

    ds = xl(make_xlsx(fill))
    assert ds.frame["x"].tolist() == [None]


def test_xlsx_hidden_sheets_are_skipped_by_default_and_not_selectable():
    data = make_xlsx(simple, hidden_sheet=True)
    assert xl(data).columns[0] == "name"
    with pytest.raises(DatasetError, match="hidden"):
        xl(data, sheet="Secret")
    with pytest.raises(DatasetError, match="not found"):
        xl(data, sheet="Nope")
    assert xl(data, sheet="Data").rows == 2


def test_xlsx_second_visible_sheet_can_be_selected():
    def fill(sheet, book):
        simple(sheet, book)
        other = book.create_sheet("Two")
        other.append(["k"])
        other.append([9])

    ds = xl(make_xlsx(fill), sheet="Two")
    assert ds.columns == ["k"] and ds.frame.loc[1, "k"] == 9


def test_xlsx_merged_and_hidden_structure_is_reported():
    def fill(sheet, _book):
        sheet.append(["a", "b", "c"])
        sheet.append([1, 2, 3])
        sheet.append([4, 5, 6])
        sheet.merge_cells("A2:B2")
        sheet.row_dimensions[3].hidden = True
        sheet.column_dimensions["C"].hidden = True

    ds = xl(make_xlsx(fill))
    assert note(ds, "merged_ranges")["count"] == 1
    assert note(ds, "hidden_rows")["count"] == 1
    assert note(ds, "hidden_columns")["count"] == 1
    assert ds.frame.loc[1, "b"] is None  # only the top-left merged cell holds a value
    assert ds.frame.loc[2, "c"] == 6  # hidden rows and columns are still read


def test_xlsx_blank_rows_and_leading_blank_rows():
    def fill(sheet, _book):
        sheet.append([])
        sheet.append(["a"])
        sheet.append([1])
        sheet.append([])
        sheet.append([3])

    ds = xl(make_xlsx(fill))
    assert ds.frame["a"].tolist() == [1, 3] and ds.frame.index.tolist() == [1, 3]


def test_xlsx_value_below_a_missing_header_is_an_error_unless_renaming():
    def fill(sheet, _book):
        sheet.append(["a"])
        sheet.append([1, 2])

    with pytest.raises(ParseError, match="empty column names"):
        xl(make_xlsx(fill))
    assert xl(make_xlsx(fill), rename_duplicates=True).columns == ["a", "column_2"]


def test_xlsx_empty_sheet_is_an_error():
    with pytest.raises(ParseError, match="empty"):
        xl(make_xlsx(lambda sheet, book: None))


def test_xlsx_forged_dimension_does_not_make_the_reader_pad_millions_of_cells():
    data = make_xlsx(simple)

    def forge(name, payload):
        if name == "xl/worksheets/sheet1.xml":
            return re.sub(rb'<dimension ref="[^"]*"', b'<dimension ref="A1:XFD1048576"', payload)
        return payload

    started = time.monotonic()
    ds = xl(rewrite(data, forge))
    assert ds.rows == 2 and time.monotonic() - started < 10


def test_xlsx_macros_nested_archives_and_old_formats_are_rejected():
    data = make_xlsx(simple)
    with_macro = rewrite(data, lambda n, p: p)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(with_macro)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info.filename))
        dst.writestr("xl/vbaProject.bin", b"x")
    with pytest.raises(DatasetError, match="macros"):
        xl(out.getvalue())

    nested = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(nested, "w") as dst:
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info.filename))
        dst.writestr("xl/embeddings/inner.zip", b"PK")
    with pytest.raises(DatasetError, match="nested archive"):
        xl(nested.getvalue())

    with pytest.raises(ParseError, match="not an .xlsx"):
        xl(b"\xd0\xcf\x11\xe0 old binary excel")


def test_xlsx_corrupt_archives_are_parse_errors_not_crashes():
    data = make_xlsx(simple)
    with pytest.raises(ParseError):
        xl(data[: len(data) // 2])
    with pytest.raises(ParseError):
        xl(b"PK\x03\x04" + b"garbage" * 10)


def bomb(size: int, name: str = "xl/worksheets/big.bin") -> bytes:
    data = make_xlsx(simple)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(data)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            dst.writestr(info.filename, src.read(info.filename))
        dst.writestr(name, b"\x00" * size)
    return out.getvalue()


def test_zip_bomb_is_stopped_by_the_ratio_and_total_size_limits():
    data = bomb(3_000_000)
    assert len(data) < 100_000  # a small archive that expands enormously
    with pytest.raises(LimitExceeded) as info:
        open_dataset(data, "t.xlsx", limits=L())
    assert info.value.limit == "zip_max_ratio"
    with pytest.raises(LimitExceeded) as info:
        open_dataset(data, "t.xlsx", limits=L(zip_max_uncompressed=1_000_000, zip_max_ratio=10**6))
    assert info.value.limit == "zip_max_uncompressed"


def test_zip_member_count_limit():
    with pytest.raises(LimitExceeded) as info:
        open_dataset(make_xlsx(simple), "t.xlsx", limits=L(zip_max_members=3))
    assert info.value.limit == "zip_max_members"


def test_xlsx_row_and_column_limits_are_counted_while_reading():
    def many(sheet, _book):
        sheet.append(["a"])
        for i in range(5):
            sheet.append([i])

    with pytest.raises(LimitExceeded) as info:
        open_dataset(make_xlsx(many), "t.xlsx", limits=L(max_rows=3))
    assert info.value.limit == "max_rows"

    def wide(sheet, _book):
        sheet.append(list("abcdef"))
        sheet.append(list(range(6)))

    with pytest.raises(LimitExceeded) as info:
        open_dataset(make_xlsx(wide), "t.xlsx", limits=L(max_columns=3))
    assert info.value.limit == "max_columns"
    with pytest.raises(LimitExceeded) as info:
        open_dataset(make_xlsx(wide), "t.xlsx", limits=L(max_cells=5))
    assert info.value.limit == "max_cells"


# --- from_frame and determinism ---------------------------------------------------------------------


def test_from_frame_wraps_a_dataframe_with_a_canonical_hash():
    frame = pd.DataFrame({"a": [1, None], "b": ["x", "y"]}, index=[5, 9])
    ds = from_frame(frame, "fetched table", kind="fetched", url="https://example.org/t")
    assert ds.frame.index.tolist() == [1, 2] and ds.frame.index.name == "row_id"
    assert str(ds.frame["a"].dtype) == "object" and ds.frame.loc[2, "a"] is None
    assert ds.source.kind == "fetched" and ds.source.url == "https://example.org/t"
    assert len(ds.source.sha256) == 64
    assert from_frame(frame, "x").source.sha256 == ds.source.sha256
    assert from_frame(frame.assign(a=[1, 2]), "x").source.sha256 != ds.source.sha256
    assert frame.index.tolist() == [5, 9]  # the caller's frame is not changed


def test_from_frame_rejects_bad_input():
    with pytest.raises(DatasetError):
        from_frame([1, 2], "x")  # type: ignore[arg-type]
    with pytest.raises(DatasetError, match="unique"):
        from_frame(pd.DataFrame([[1, 2]], columns=["a", "a"]), "x")


def test_loading_the_same_bytes_twice_gives_identical_tables():
    data = b"a,b\n1,x\n2,y\n"
    first, second = csv_ds(data), csv_ds(data)
    assert first.frame.equals(second.frame) and first.source == second.source
