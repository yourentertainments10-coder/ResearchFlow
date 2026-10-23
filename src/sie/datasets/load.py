"""Track A: read bytes into a ``Dataset`` (CSV, XLSX, JSON) with limits. See design sections 8 and 10.

Nothing here touches the file system or the network: the caller has read the file once and hands over
the bytes. Every policy either follows the design or raises a specific error that names the place.
A partial dataset is never returned.
"""

from __future__ import annotations

import csv
import io
import json
import math
import time
import zipfile
from typing import Any
from xml.etree import ElementTree

import pandas as pd

from sie.datasets.models import (
    Dataset,
    DatasetError,
    LimitExceeded,
    Limits,
    LoadOptions,
    ParseError,
    SourceRef,
)
from sie.datasets.provenance import sha256_bytes, table_sha256

CHECK_EVERY_ROWS = 10_000
_REJECT_COLUMNS = ["row_id", "operation", "column", "original_value", "reason"]
_UNSUPPORTED_EXCEL = (".xls", ".xlsb", ".xlsm")
_NESTED_ARCHIVE_SUFFIXES = (
    ".zip",
    ".xlsx",
    ".xlsm",
    ".docx",
    ".pptx",
    ".jar",
    ".gz",
    ".7z",
    ".rar",
    ".tar",
)
_CHUNK = 64 * 1024
_INT64 = 2**63


class _Budget:
    """Wall-clock budget, checked every ``CHECK_EVERY_ROWS`` rows (design section 8)."""

    def __init__(self, limits: Limits) -> None:
        self._limit = limits.parse_seconds
        self._start = time.monotonic()

    def tick(self, rows_done: int) -> None:
        if rows_done % CHECK_EVERY_ROWS == 0 and time.monotonic() - self._start > self._limit:
            raise LimitExceeded(
                "parse_seconds", f"parsing took longer than {self._limit:g} seconds"
            )


def _note(code: str, count: int, detail: Any = None) -> dict[str, Any]:
    entry: dict[str, Any] = {"code": code, "count": count}
    if detail is not None:
        entry["detail"] = detail
    return entry


def _frame(columns: list[str], rows: list[list[Any]], row_ids: list[int]) -> pd.DataFrame:
    frame = pd.DataFrame(rows, columns=columns, dtype=object)
    if rows:
        frame = frame.astype(object)
    frame.index = pd.Index(row_ids, name="row_id", dtype="int64")
    return frame


def _reject_frame(records: list[tuple[int, str, str]]) -> pd.DataFrame:
    return pd.DataFrame(
        [(rid, "load", None, raw, reason) for rid, raw, reason in records],
        columns=_REJECT_COLUMNS,
        dtype=object,
    )


# --- headers (shared by CSV, Excel and JSON) -------------------------------------------------------


def _headers(
    raw: list[Any], options: LoadOptions, limits: Limits
) -> tuple[list[str], list[dict[str, Any]]]:
    """Trim names, then apply the empty/duplicate policy: an error naming positions, or renaming."""
    if len(raw) > limits.max_columns:
        raise LimitExceeded("max_columns", f"{len(raw)} columns, limit {limits.max_columns}")
    trimmed = ["" if value is None else str(value).strip() for value in raw]
    for position, name in enumerate(trimmed, start=1):
        if len(name) > limits.max_cell_chars:
            raise LimitExceeded("max_cell_chars", f"header {position} is longer than the limit")
    empty = [i for i, name in enumerate(trimmed, start=1) if name == ""]
    positions: dict[str, list[int]] = {}
    for i, name in enumerate(trimmed, start=1):
        if name:
            positions.setdefault(name, []).append(i)
    duplicates = {name: where for name, where in positions.items() if len(where) > 1}
    if (empty or duplicates) and not options.rename_duplicates:
        parts = []
        if empty:
            parts.append(f"empty column names at positions {empty}")
        if duplicates:
            listed = "; ".join(
                f"{name!r} at positions {where}" for name, where in duplicates.items()
            )
            parts.append(f"duplicate column names: {listed}")
        raise ParseError(
            "; ".join(parts) + " (use rename_duplicates / --rename-duplicates to rename them)",
            "header",
        )
    reserved = set(positions)
    used: set[str] = set()
    names: list[str] = []
    renames: list[dict[str, Any]] = []
    for i, name in enumerate(trimmed, start=1):
        final = name
        if name == "":
            final = f"column_{i}"
            while final in reserved or final in used:
                final += "_"
        elif name in used:
            k = 2
            while f"{name}__{k}" in reserved or f"{name}__{k}" in used:
                k += 1
            final = f"{name}__{k}"
        if final != name:
            renames.append({"position": i, "from": name, "to": final})
        used.add(final)
        names.append(final)
    notes = [_note("renamed_header", len(renames), renames)] if renames else []
    return names, notes


# --- CSV ---------------------------------------------------------------------------------------------


def _decode(data: bytes, options: LoadOptions) -> str:
    encoding = options.encoding.lower().replace("_", "-")
    if encoding in ("utf-8", "utf8", "utf-8-sig", "utf8-sig"):
        if data.startswith(b"\xef\xbb\xbf"):
            data = data[3:]
        codec = "utf-8"
    else:
        codec = options.encoding
    try:
        text = data.decode(codec, errors="strict")
    except UnicodeDecodeError as exc:
        raise ParseError(
            f"bytes are not valid {codec}; if the file uses another encoding pass it explicitly",
            f"byte offset {exc.start}",
        ) from exc
    except LookupError as exc:
        raise DatasetError(f"unknown encoding {options.encoding!r}") from exc
    if "\x00" in text:
        raise ParseError("NUL byte in file", f"character offset {text.index(chr(0))}")
    return text


def _load_csv(
    data: bytes, options: LoadOptions, limits: Limits
) -> tuple[pd.DataFrame, pd.DataFrame | None, list[dict[str, Any]]]:
    text = _decode(data, options)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=options.delimiter, strict=True)
    budget = _Budget(limits)
    notes: list[dict[str, Any]] = []

    def next_record() -> list[str] | None:
        try:
            return next(reader)
        except StopIteration:
            return None
        except csv.Error as exc:
            if "field larger than field limit" in str(exc):
                raise LimitExceeded(
                    "max_cell_chars", "a field is larger than the parser allows"
                ) from exc
            raise ParseError(f"malformed CSV: {exc}", f"line {reader.line_num}") from exc

    header = next_record()
    while header is not None and header == []:
        header = next_record()  # leading blank lines carry no information
    if header is None:
        raise ParseError("file is empty: no header row")
    if len(header) == 1 and options.delimiter == ",":
        for other, label in ((";", "semicolon"), ("\t", "tab")):
            if other in header[0]:
                raise ParseError(
                    f"the header is a single column that contains a {label}; "
                    f"the delimiter is not guessed, pass delimiter={other!r} (--delimiter)",
                    "header",
                )
    names, header_notes = _headers(list(header), options, limits)
    notes.extend(header_notes)
    width = len(names)

    rows: list[list[Any]] = []
    row_ids: list[int] = []
    rejected: list[tuple[int, str, str]] = []
    blanks: list[int] = []
    record_no = 0
    cells = 0
    while True:
        fields = next_record()
        if fields is None:
            break
        record_no += 1
        budget.tick(record_no)
        if fields == []:
            blanks.append(record_no)
            continue
        if len(fields) != width:
            if options.ragged == "reject":
                rejected.append((record_no, options.delimiter.join(fields), "RAGGED_ROW"))
                continue
            raise ParseError(
                f"row has {len(fields)} fields but the header has {width}; "
                "pass ragged='reject' (--ragged reject) to set such rows aside",
                f"row {record_no}",
            )
        if len(rows) + 1 > limits.max_rows:
            raise LimitExceeded("max_rows", f"more than {limits.max_rows} rows")
        cells += width
        if cells > limits.max_cells:
            raise LimitExceeded("max_cells", f"more than {limits.max_cells} cells")
        for value in fields:
            if len(value) > limits.max_cell_chars:
                raise LimitExceeded(
                    "max_cell_chars", f"a cell in row {record_no} is longer than the limit"
                )
        rows.append([None if value == "" else value for value in fields])
        row_ids.append(record_no)
    if blanks:
        notes.append(_note("blank_lines_skipped", len(blanks), {"row_ids": blanks[:20]}))
    rejects = _reject_frame(rejected) if rejected else None
    if rejected:
        notes.append(_note("ragged_rows_rejected", len(rejected)))
    return _frame(names, rows, row_ids), rejects, notes


# --- Excel -------------------------------------------------------------------------------------------


def _check_archive(data: bytes, limits: Limits) -> tuple[zipfile.ZipFile, list[dict[str, Any]]]:
    """ZIP-bomb defence: declared and actual sizes, ratios, member count, nested archives, macros."""
    if not data.startswith(b"PK"):
        raise ParseError("not an .xlsx file (it is not a ZIP archive; .xls is not supported)")
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as exc:
        raise ParseError(f"not a valid .xlsx archive: {exc}") from exc
    infos = archive.infolist()
    if len(infos) > limits.zip_max_members:
        raise LimitExceeded(
            "zip_max_members", f"{len(infos)} members, limit {limits.zip_max_members}"
        )
    names = [info.filename for info in infos]
    lowered = [name.lower() for name in names]
    if any(name.endswith("vbaproject.bin") for name in lowered):
        raise DatasetError("workbook contains macros (vbaProject.bin); it is rejected")
    for name in lowered:
        if name.endswith(_NESTED_ARCHIVE_SUFFIXES):
            raise DatasetError(f"workbook contains a nested archive ({name}); it is rejected")
    declared_total = sum(info.file_size for info in infos)
    if declared_total > limits.zip_max_uncompressed:
        raise LimitExceeded(
            "zip_max_uncompressed",
            f"declared uncompressed size {declared_total} exceeds {limits.zip_max_uncompressed}",
        )
    for info in infos:
        if (
            info.file_size > limits.zip_ratio_min_bytes
            and info.file_size > limits.zip_max_ratio * max(info.compress_size, 1)
        ):
            raise LimitExceeded(
                "zip_max_ratio",
                f"member {info.filename!r} compresses more than {limits.zip_max_ratio}:1",
            )
    # declared sizes can be forged: decompress in chunks with a running counter and stop at the cap
    total = 0
    for info in infos:
        member = 0
        with archive.open(info) as handle:
            while True:
                try:
                    chunk = handle.read(_CHUNK)
                except (
                    zipfile.BadZipFile,
                    EOFError,
                    OSError,
                    RuntimeError,
                    NotImplementedError,
                ) as exc:
                    raise ParseError(f"corrupt archive member {info.filename!r}: {exc}") from exc
                if not chunk:
                    break
                member += len(chunk)
                total += len(chunk)
                if total > limits.zip_max_uncompressed:
                    raise LimitExceeded(
                        "zip_max_uncompressed",
                        f"archive expands beyond {limits.zip_max_uncompressed} bytes",
                    )
                if member > limits.zip_ratio_min_bytes and member > limits.zip_max_ratio * max(
                    info.compress_size, 1
                ):
                    raise LimitExceeded(
                        "zip_max_ratio",
                        f"member {info.filename!r} compresses more than {limits.zip_max_ratio}:1",
                    )
    notes: list[dict[str, Any]] = []
    links = [n for n in names if n.lower().startswith("xl/externallinks/")]
    objects = [n for n in names if n.lower().startswith(("xl/embeddings/", "xl/activex/"))]
    if links:
        notes.append(_note("external_links_ignored", len(links)))
    if objects:
        notes.append(_note("embedded_objects_ignored", len(objects)))
    return archive, notes


def _sheet_xml_counts(archive: zipfile.ZipFile, part: str | None) -> dict[str, int] | None:
    """Merged ranges and hidden rows/columns of one sheet, streamed from its XML part."""
    if not part:
        return None
    path = part.lstrip("/")
    if path not in archive.namelist():
        path = "xl/" + path
    if path not in archive.namelist():
        return None
    counts = {"merged_ranges": 0, "hidden_rows": 0, "hidden_columns": 0}
    try:
        with archive.open(path) as handle:
            for _event, element in ElementTree.iterparse(handle, events=("start",)):
                tag = element.tag.rsplit("}", 1)[-1]
                if tag == "mergeCell":
                    counts["merged_ranges"] += 1
                elif tag in ("row", "col") and element.get("hidden") in ("1", "true"):
                    counts["hidden_rows" if tag == "row" else "hidden_columns"] += 1
                element.clear()
    except ElementTree.ParseError:
        return None
    return counts


def _trim_row(values: list[Any]) -> list[Any]:
    end = len(values)
    while end and values[end - 1] is None:
        end -= 1
    return values[:end]


def _load_xlsx(
    data: bytes, options: LoadOptions, limits: Limits
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    from openpyxl import load_workbook

    archive, notes = _check_archive(data, limits)
    budget = _Budget(limits)
    try:
        values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formula_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except DatasetError:
        raise
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many types for corrupt files
        raise ParseError(f"cannot read workbook: {type(exc).__name__}: {exc}") from exc
    try:
        visible = [s.title for s in values_book.worksheets if s.sheet_state == "visible"]
        if options.sheet is None:
            if not visible:
                raise DatasetError("workbook has no visible sheet")
            title = visible[0]
        else:
            all_titles = [s.title for s in values_book.worksheets]
            if options.sheet not in all_titles:
                raise DatasetError(f"sheet {options.sheet!r} not found; visible sheets: {visible}")
            if options.sheet not in visible:
                raise DatasetError(f"sheet {options.sheet!r} is hidden and cannot be selected")
            title = options.sheet
        sheet = values_book[title]
        formula_sheet = formula_book[title]
        # a forged <dimension> must not make the reader pad millions of empty cells
        sheet.reset_dimensions()
        formula_sheet.reset_dimensions()
        epoch = getattr(values_book, "epoch", None)
        # openpyxl reports 1899-12-30 for the 1900 date system and 1904-01-01 for the 1904 system
        epoch_system = 1904 if getattr(epoch, "year", None) == 1904 else 1900
        notes.append(_note("excel_epoch", epoch_system, {"sheet": title}))

        raw_rows: list[list[Any]] = []
        formulas = 0
        without_cache = 0
        error_cells = 0
        cells = 0
        for row_no, (values_row, formula_row) in enumerate(
            zip(sheet.iter_rows(), formula_sheet.iter_rows(), strict=False), start=1
        ):
            budget.tick(row_no)
            if row_no > limits.max_rows + 1:
                raise LimitExceeded("max_rows", f"more than {limits.max_rows} rows")
            values: list[Any] = []
            row_has_formula = False
            for value_cell, formula_cell in zip(values_row, formula_row, strict=False):
                value = value_cell.value
                if getattr(formula_cell, "data_type", None) == "f":
                    formulas += 1
                    row_has_formula = True
                    if value is None:
                        without_cache += 1
                elif getattr(value_cell, "data_type", None) == "e":
                    error_cells += 1
                if isinstance(value, str) and len(value) > limits.max_cell_chars:
                    raise LimitExceeded(
                        "max_cell_chars", f"a cell in sheet row {row_no} is longer than the limit"
                    )
                values.append(value)
            trimmed = _trim_row(values)
            if row_has_formula and not trimmed:
                trimmed = [None]  # a row holding only formulas without cached values is not blank
            if len(trimmed) > limits.max_columns:
                raise LimitExceeded(
                    "max_columns",
                    f"row {row_no} uses {len(trimmed)} columns, limit {limits.max_columns}",
                )
            cells += len(trimmed)
            if cells > limits.max_cells:
                raise LimitExceeded("max_cells", f"more than {limits.max_cells} cells")
            raw_rows.append(trimmed)
        part = getattr(sheet, "_worksheet_path", None)
    finally:
        values_book.close()
        formula_book.close()

    while raw_rows and not raw_rows[0]:
        raw_rows.pop(0)  # blank rows above the header
    if not raw_rows:
        raise ParseError("sheet is empty: no header row", f"sheet {title!r}")
    header_row = raw_rows[0]
    body = raw_rows[1:]
    width = max([len(header_row)] + [len(r) for r in body])
    names, header_notes = _headers(header_row + [None] * (width - len(header_row)), options, limits)
    notes.extend(header_notes)
    rows: list[list[Any]] = []
    row_ids: list[int] = []
    blanks = 0
    for record_no, values in enumerate(body, start=1):
        if not values:
            blanks += 1
            continue
        rows.append(values + [None] * (width - len(values)))
        row_ids.append(record_no)
    if len(rows) > limits.max_rows:
        raise LimitExceeded("max_rows", f"{len(rows)} rows, limit {limits.max_rows}")
    # drop blank rows at the very end silently counted; blanks between rows are counted too
    if blanks:
        notes.append(_note("blank_rows_skipped", blanks))
    notes.append(_note("formula_cells", formulas, {"never_evaluated": True}))
    if without_cache:
        notes.append(_note("formula_without_cached_value", without_cache))
    if error_cells:
        notes.append(_note("excel_error_cells", error_cells))
    counts = _sheet_xml_counts(archive, part)
    if counts is None:
        notes.append(
            _note("sheet_structure_unavailable", 0, "merged and hidden counts could not be read")
        )
    else:
        for key, value in counts.items():
            if value:
                notes.append(_note(key, value))
    return _frame(names, rows, row_ids), notes


# --- JSON --------------------------------------------------------------------------------------------


def _depth_of(value: Any, limit: int) -> int:
    deepest = 0
    stack: list[tuple[Any, int]] = [(value, 1)]
    while stack:
        item, depth = stack.pop()
        deepest = max(deepest, depth)
        if deepest > limit:
            return deepest
        if isinstance(item, dict):
            stack.extend(
                (child, depth + 1) for child in item.values() if isinstance(child, dict | list)
            )
        elif isinstance(item, list):
            stack.extend((child, depth + 1) for child in item if isinstance(child, dict | list))
    return deepest


def _strict_pairs(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    seen: dict[str, Any] = {}
    for key, value in pairs:
        if key in seen:
            raise ParseError(f"duplicate key {key!r} in an object")
        seen[key] = value
    return seen


def _reject_constant(token: str) -> Any:
    raise ParseError(f"{token} is not allowed in JSON data")


def _parse_int(token: str) -> int:
    number = int(token)
    if not -_INT64 <= number < _INT64:
        raise ParseError(f"integer {token[:30]} is outside the 64-bit range")
    return number


def _parse_float(token: str) -> float:
    number = float(token)
    if not math.isfinite(number):
        raise ParseError(f"number {token[:30]} is not finite")
    return number


def _load_json(
    data: bytes, options: LoadOptions, limits: Limits
) -> tuple[pd.DataFrame, list[dict[str, Any]]]:
    text = _decode(data, LoadOptions(encoding="utf-8"))
    try:
        document = json.loads(
            text,
            object_pairs_hook=_strict_pairs,
            parse_constant=_reject_constant,
            parse_int=_parse_int,
            parse_float=_parse_float,
        )
    except json.JSONDecodeError as exc:
        raise ParseError(
            f"malformed JSON: {exc.msg}", f"line {exc.lineno}, column {exc.colno}"
        ) from exc
    except RecursionError as exc:
        raise LimitExceeded(
            "max_json_depth", f"nesting is deeper than {limits.max_json_depth}"
        ) from exc
    except ValueError as exc:
        if isinstance(exc, DatasetError):
            raise
        raise ParseError(f"cannot read JSON: {exc}") from exc
    if not isinstance(document, list):
        raise ParseError("JSON must be an array of objects", f"found {type(document).__name__}")
    if not document:
        raise ParseError("JSON array is empty: no columns to read")
    if len(document) > limits.max_rows:
        raise LimitExceeded("max_rows", f"{len(document)} rows, limit {limits.max_rows}")
    budget = _Budget(limits)
    order: list[str] = []
    seen: set[str] = set()
    for index, item in enumerate(document, start=1):
        budget.tick(index)
        if not isinstance(item, dict):
            raise ParseError(
                "every array element must be an object", f"row {index}: {type(item).__name__}"
            )
        for key in item:
            if key not in seen:
                seen.add(key)
                order.append(key)
    names, notes = _headers(order, options, limits)
    mapping = dict(zip(order, names, strict=True))
    if len(order) * len(document) > limits.max_cells:
        raise LimitExceeded("max_cells", f"more than {limits.max_cells} cells")
    rows: list[list[Any]] = []
    missing: dict[str, int] = {}
    for index, item in enumerate(document, start=1):
        budget.tick(index)
        row: list[Any] = []
        for key in order:
            if key not in item:
                missing[mapping[key]] = missing.get(mapping[key], 0) + 1
                row.append(None)
                continue
            value = item[key]
            if isinstance(value, dict | list):
                if _depth_of(value, limits.max_json_depth) > limits.max_json_depth:
                    raise LimitExceeded(
                        "max_json_depth", f"nesting is deeper than {limits.max_json_depth}"
                    )
                kind = "object" if isinstance(value, dict) else "array"
                raise ParseError(
                    f"nested {kind} is not supported (flattening would be a guess)",
                    f"row {index}, path ${'[' + str(index - 1) + ']'}.{key}",
                )
            if isinstance(value, str) and len(value) > limits.max_cell_chars:
                raise LimitExceeded(
                    "max_cell_chars", f"a value in row {index} is longer than the limit"
                )
            row.append(value)
        rows.append(row)
    if missing:
        notes.append(_note("missing_keys", sum(missing.values()), missing))
    return _frame(names, rows, list(range(1, len(rows) + 1))), notes


# --- public API --------------------------------------------------------------------------------------


def _extension(name: str) -> str:
    lowered = name.lower()
    dot = lowered.rfind(".")
    return lowered[dot:] if dot >= 0 else ""


def open_dataset(
    data: bytes,
    name: str,
    options: LoadOptions | None = None,
    limits: Limits | None = None,
    *,
    kind: str = "upload",
    url: str | None = None,
    retrieved_at: str | None = None,
) -> Dataset:
    """Detect the format from ``name`` (.csv, .xlsx, .json) and parse ``data``.

    The caller has already read the file once; this function never touches the file system.
    Raises ``LimitExceeded``, ``ParseError`` or ``DatasetError``; never returns a partial dataset.
    """
    opts = options or LoadOptions()
    lim = limits or Limits()
    if not isinstance(data, bytes | bytearray):
        raise DatasetError("data must be bytes")
    data = bytes(data)
    if len(data) > lim.max_bytes:
        raise LimitExceeded("max_bytes", f"{len(data)} bytes, limit {lim.max_bytes}")
    extension = _extension(name)
    rejects: pd.DataFrame | None = None
    if extension == ".csv":
        frame, rejects, notes = _load_csv(data, opts, lim)
    elif extension == ".xlsx":
        frame, notes = _load_xlsx(data, opts, lim)
    elif extension == ".json":
        frame, notes = _load_json(data, opts, lim)
    elif extension in _UNSUPPORTED_EXCEL:
        raise DatasetError(f"{extension} files are not supported; save the sheet as .xlsx or .csv")
    else:
        raise DatasetError(
            f"unsupported file type {extension or '(none)'!r}; use .csv, .xlsx or .json"
        )
    source = SourceRef(
        kind=kind,  # type: ignore[arg-type]
        name=name,
        sha256=sha256_bytes(data),
        size_bytes=len(data),
        url=url,
        retrieved_at=retrieved_at,
    )
    return Dataset(frame=frame, source=source, options=opts, rejects=rejects, notes=notes)


def from_frame(
    frame: pd.DataFrame,
    name: str,
    *,
    kind: str = "derived",
    url: str | None = None,
    retrieved_at: str | None = None,
) -> Dataset:
    """Wrap a DataFrame a caller already holds (for example a future fetcher). Index becomes row_id
    1..n, columns become object dtype, ``SourceRef.sha256`` is the canonical table hash."""
    if not isinstance(frame, pd.DataFrame):
        raise DatasetError("from_frame needs a pandas DataFrame")
    names = [str(c) for c in frame.columns]
    if len(set(names)) != len(names):
        raise DatasetError("column names must be unique")
    copy = frame.copy()
    copy.columns = names
    copy = copy.astype(object)
    copy = copy.where(copy.notna(), None)
    copy.index = pd.Index(range(1, len(copy) + 1), name="row_id", dtype="int64")
    source = SourceRef(
        kind=kind,  # type: ignore[arg-type]
        name=name,
        sha256=table_sha256(copy),
        size_bytes=0,
        url=url,
        retrieved_at=retrieved_at,
    )
    return Dataset(frame=copy, source=source)


__all__ = ["from_frame", "open_dataset"]
